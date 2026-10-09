import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Theme Colors (Modern Minimalist / Slate & Indigo)
    COLOR_BG = RGBColor(248, 250, 252)        # Slate 50 (#F8FAFC)
    COLOR_CARD = RGBColor(255, 255, 255)      # White
    COLOR_BORDER = RGBColor(226, 232, 240)    # Slate 200 (#E2E8F0)
    COLOR_TEXT_MAIN = RGBColor(15, 23, 42)    # Slate 900 (#0F172A)
    COLOR_TEXT_MUTED = RGBColor(71, 85, 105)  # Slate 600 (#475569)
    COLOR_ACCENT = RGBColor(79, 70, 229)      # Indigo 600 (#4F46E5)
    COLOR_ACCENT_BG = RGBColor(238, 242, 255) # Indigo 50 (#EEF2FF)
    COLOR_TAG_TEXT = RGBColor(67, 56, 202)    # Indigo 700 (#4338CA)

    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text=""):
        # Tag pill
        if tag_text:
            tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(4.5), Inches(0.4))
            tf_tag = tag_box.text_frame
            tf_tag.word_wrap = True
            p_tag = tf_tag.paragraphs[0]
            p_tag.text = tag_text.upper()
            p_tag.font.size = Pt(10)
            p_tag.font.bold = True
            p_tag.font.color.rgb = COLOR_ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_MAIN

        # Subtitle
        if subtitle_text:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.5))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle_text
            p_sub.font.size = Pt(12)
            p_sub.font.color.rgb = COLOR_TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, badge=""):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        tx_box = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        if badge:
            p0.text = f"{badge}  •  {title}"
        else:
            p0.text = title
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_TEXT_MAIN

        for it in items:
            p = tf.add_paragraph()
            p.text = f"•  {it}"
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_TEXT_MUTED
            p.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 1 : COVER
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s1)

    # Big Title Box
    title_box = s1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(3.5))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "POINT D'ÉTAPE PÉDAGOGIQUE — WEMODO"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_ACCENT
    p_badge.space_after = Pt(14)

    p_main = tf1.add_paragraph()
    p_main.text = "Architecture des Modules 4 & 5"
    p_main.font.size = Pt(38)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_TEXT_MAIN
    p_main.space_after = Pt(16)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Validation de la structure, des pistes de projet fil rouge et des choix d'arbitrage avant tournage"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom footer card
    card_info = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.6), Inches(10.9), Inches(0.9))
    card_info.fill.solid()
    card_info.fill.fore_color.rgb = COLOR_CARD
    card_info.line.color.rgb = COLOR_BORDER
    card_info.line.width = Pt(1)

    tx_info = s1.shapes.add_textbox(Inches(1.5), Inches(5.75), Inches(10.3), Inches(0.6))
    tf_info = tx_info.text_frame
    p_info = tf_info.paragraphs[0]
    p_info.text = "Objectif de la séance : Aligner l'équipe sur la séparation « Boîte à outils (M4) » vs « Sprint Certifiant (M5) » et arrêter le projet de fin d'études."
    p_info.font.size = Pt(12)
    p_info.font.color.rgb = COLOR_TEXT_MAIN

    # -------------------------------------------------------------
    # SLIDE 2 : VISION & PHILOSOPHIE
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s2)
    add_header(s2, "Philosophie Pédagogique", "La nouvelle articulation : Apprentissage vs Exécution",
               "Séparer l'acquisition des briques techniques et la réalisation du projet certifiant pour maximiser la réussite.")

    add_card(s2, 0.8, 2.0, 5.7, 4.8, "MODULE 4 : L'Armurerie & la Boîte à Outils", [
        "Objectif : Découvrir, tester et maîtriser chaque technologie une par une.",
        "Projet support : Application d'exercice (ex: Wander) pour ancrer la pratique.",
        "Environnement local : Installation d'Antigravity et utilisation du Terminal.",
        "Persistance : Connexion Supabase (BaaS) et création de tables SQL via MCP.",
        "Sécurité & Accès : Authentification (Email/OAuth) et règles de RLS.",
        "Postulat : L'apprenant a le droit à l'erreur et expérimente sans la pression de l'examen."
    ], "M4 — APPRENDRE")

    add_card(s2, 6.8, 2.0, 5.7, 4.8, "MODULE 5 : Le Sprint Certifiant (Conditions Réelles)", [
        "Objectif : Réaliser son projet de certification de A à Z en autonomie guidée.",
        "Nouveau sujet : Projet inédit (Wander s'arrête en fin de M4).",
        "Méthodologie pro : Cadrage formalisé ➔ Dépôt Git ➔ 7 étapes chronologiques de dev.",
        "100 % conforme à la certification : Couverture stricte des 6 compétences (C1 à C6).",
        "Livrables officiels générés : Dossier PDF, Repo GitHub, App Vercel, Vidéo de démo.",
        "Postulat : L'apprenant applique les compétences acquises en M4 sans réexplication théorique."
    ], "M5 — PRODUIRE")

    # -------------------------------------------------------------
    # SLIDE 3 : MODULE 4 EN DÉTAIL
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s3)
    add_header(s3, "Contenu du Programme", "Module 4 : Antigravity, BDD & Authentification",
               "29 leçons réparties en 4 chapitres pour doter l'apprenant de tous ses super-pouvoirs techniques.")

    add_card(s3, 0.8, 2.0, 2.75, 4.8, "Ch. 1 — Bases d’Antigravity", [
        "8 leçons",
        "Pourquoi un IDE local",
        "Installation & setup",
        "Découverte interface",
        "Commandes Terminal clés",
        "Exécution de scripts",
        "Prise en main de l'agent"
    ], "1")

    add_card(s3, 3.75, 2.0, 2.75, 4.8, "Ch. 2 — Évolution de projet", [
        "7 leçons",
        "Démarrage from scratch",
        "Reprise de projet existant",
        "Variables & secrets (.env)",
        "Lancement en local",
        "Audit & itération agent",
        "Débogage pas-à-pas"
    ], "2")

    add_card(s3, 6.7, 2.0, 2.75, 4.8, "Ch. 3 — Base de données", [
        "7 leçons",
        "Découverte Supabase BaaS",
        "Création projet Cloud",
        "Connexion via le MCP",
        "Modélisation SQL par IA",
        "Connexion du client JS",
        "Persistance dans l'UI"
    ], "3")

    add_card(s3, 9.65, 2.0, 2.85, 4.8, "Ch. 4 — Auth & RLS", [
        "7 leçons",
        "Gestion session & JWT",
        "Portail Email/Password",
        "Connexion Google (OAuth)",
        "Protection routes privées",
        "Théorie de la faille RLS",
        "Politiques RLS SQL 🚨"
    ], "4")

    # -------------------------------------------------------------
    # SLIDE 4 : MODULE 5 EN DÉTAIL
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s4)
    add_header(s4, "Contenu du Programme", "Module 5 : Réussir son Projet Certifiant",
               "19 leçons en 3 chapitres — Un accompagnement pas-à-pas pour sécuriser les 4 livrables du jury.")

    add_card(s4, 0.8, 2.0, 3.65, 4.8, "Ch. 1 — Cadrer son projet", [
        "4 leçons  •  Livrable n°1",
        "Décryptage grille d'évaluation",
        "Périmètre & Persona (CE1.1)",
        "Schéma Mermaid du parcours utilisateur (CE1.2)",
        "Règles contextuelles structurantes avec AGENTS.md (CE1.3)"
    ], "CADRAGE")

    add_card(s4, 4.75, 2.0, 3.85, 4.8, "Ch. 2 — Le Sprint de Dév", [
        "8 leçons  •  7 Étapes",
        "Étape 1 : Générer le prototype initial",
        "Étape 2 : Init Git & reprise locale (CE3.1)",
        "Étape 3 : Injecter contexte AGENTS.md",
        "Étape 4 : Intégrer l'API externe",
        "Étape 5 : Connecter BDD & Auth",
        "Étape 6 : Verrouiller avec le RLS 🚨",
        "Étape 7 : Audit ergonomie & bugs"
    ], "DÉVELOPPEMENT")

    add_card(s4, 8.9, 2.0, 3.65, 4.8, "Ch. 3 — Prod & Certification", [
        "7 leçons  •  Livrables 2, 3, 4",
        "Refactoring propre (CE3.3)",
        "Dépôt GitHub + README (Livrable 3)",
        "Déploiement Vercel HTTPS (Livrable 2)",
        "Audit sécurité étanchéité données",
        "Vidéo de démo 3-5 min (Livrable 4)",
        "Préparation & conseils QCM"
    ], "LIVRAISON")

    # -------------------------------------------------------------
    # SLIDE 5 : PISTES DE PROJET FIL ROUGE POUR LE M5
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s5)
    add_header(s5, "Projet Certifiant", "Pistes de Projet Fil Rouge pour le Module 5",
               "Le projet fil rouge doit obligatoirement combiner : IA générative, BDD Cloud, Auth privée, RLS et Vercel.")

    add_card(s5, 0.8, 2.0, 3.65, 4.8, "Piste 1 : Outil Métier & B2B", [
        "Typologie : Prospection & Qualification de leads",
        "Valeur ajoutée IA : Analyse du profil de l'entreprise et rédaction de pitchs personnalisés.",
        "Base de données : Table des prospects, statuts, historique des échanges.",
        "Auth & RLS : Espace commercial privé, un commercial ne voit que ses propres leads.",
        "Atout : Profil très valorisable sur le marché du travail."
    ], "B2B / VENTE")

    add_card(s5, 4.75, 2.0, 3.85, 4.8, "Piste 2 : Plateforme Marketing", [
        "Typologie : Planificateur & Générateur de Contenu",
        "Valeur ajoutée IA : Génération de calendrier éditorial, posts réseaux sociaux & newsletters.",
        "Base de données : Bibliothèque de contenus, statuts (brouillon, validé, publié).",
        "Auth & RLS : Espace créateur / marque sécurisé, isolation complète des publications.",
        "Atout : Cas d'usage très universel, ludique et immédiat."
    ], "MARKETING")

    add_card(s5, 8.9, 2.0, 3.65, 4.8, "Piste 3 : Suivi & Coaching RH", [
        "Typologie : Dashboard d'Objectifs & Suivi Personnalisé",
        "Valeur ajoutée IA : Génération de plans d'action sur-mesure selon les forces/faiblesses.",
        "Base de données : Objectifs, étapes cochées, calcul de métriques en direct.",
        "Auth & RLS : Données personnelles sensibles avec justification parfaite du RLS.",
        "Atout : Démontre la gestion rigoureuse de données confidentielles."
    ], "RH / COACHING")

    # -------------------------------------------------------------
    # SLIDE 6 : ARBITRAGES PÉDAGOGIQUES
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s6)
    add_header(s6, "Points de Décision", "4 Arbitrages Clés pour l'Équipe Pédagogique",
               "Les choix stratégiques à trancher ensemble pour sécuriser la qualité pédagogique et le tournage.")

    add_card(s6, 0.8, 2.0, 5.7, 2.3, "1. Choix du Sujet du Module 5", [
        "Quelle piste thématique retenons-nous pour le tournage (B2B, Marketing, ou Coaching) ?",
        "L'apprenant suit-il le même sujet ou est-il encouragé à personnaliser son propre projet ?"
    ], "ARBITRAGE 1")

    add_card(s6, 6.8, 2.0, 5.7, 2.3, "2. Workflow Git ➔ Vercel", [
        "On valide l'initialisation de Git dès le début du Chapitre 2 (CE3.1 : commits réguliers).",
        "Déploiement Vercel connecté au repo GitHub (meilleure pratique pro et automatisée)."
    ], "ARBITRAGE 2")

    add_card(s6, 0.8, 4.5, 5.7, 2.3, "3. Fonctions Avancées (Skills / Agents)", [
        "Intégration comme accélérateurs dans les 7 étapes (Browser Agent pour l'UI, sous-agent refactoring).",
        "Évite d'alourdir le programme tout en montrant la puissance d'Antigravity."
    ], "ARBITRAGE 3")

    add_card(s6, 6.8, 4.5, 5.7, 2.3, "4. Sécurisation du Dossier de Cadrage", [
        "Fournir un modèle de dossier (Markdown/Notion) pour guider l'apprenant.",
        "Intégrer le schéma Mermaid dès l'étape de cadrage pour valider le critère CE1.2."
    ], "ARBITRAGE 4")

    # -------------------------------------------------------------
    # SLIDE 7 : ROADMAP & NEXT STEPS
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s7)
    add_header(s7, "Planning de Production", "Roadmap & Prochaines Étapes",
               "Calendrier des actions pour finaliser les modules 4 et 5 avant la date limite de mise en ligne.")

    add_card(s7, 0.8, 2.0, 3.65, 4.8, "Étape 1 : Cadrage & Scripts", [
        "Validation finale des arbitrages aujourd'hui",
        "Rédaction des scripts des leçons théoriques",
        "Création des prompts types pour l'agent IA",
        "Préparation du code starter du projet M5",
        "Durée estimée : 3 à 5 jours"
    ], "JUSQU'À J+5")

    add_card(s7, 4.75, 2.0, 3.85, 4.8, "Étape 2 : Tournage & Démos", [
        "Enregistrement des leçons logicielles (screencasts)",
        "Démonstrations des 7 étapes de développement",
        "Vérification en direct du RLS et du déploiement Vercel",
        "Enregistrement des capsules théoriques",
        "Durée estimée : 4 à 6 jours"
    ], "J+5 À J+10")

    add_card(s7, 8.9, 2.0, 3.65, 4.8, "Étape 3 : Recette & Sortie", [
        "Montage vidéo et habillage graphique",
        "Vérification complète de la grille certifiante",
        "Génération des quiz d'auto-évaluation (QCM)",
        "Publication sur la plateforme Wemodo",
        "Durée estimée : 3 jours"
    ], "J+10 À J+15")

    output_path = "/Users/maximeelhaik/Documents/VIBE CODING GENERATION/Presentation_Pedagogique_Wemodo_M4_M5.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
